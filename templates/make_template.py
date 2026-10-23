"""make_template.py - builds the Word template templates/contract_template.docx.

What this file is for:
    It rebuilds the law student's contract skeleton (Construction_Contract_Skeleton_EN)
    word for word as a Word file with gaps for docxtpl. The gaps use double braces,
    e.g. {{ client_name }} or {{ price_clause }}; generator.py fills them (step 15).
    Lines like {%p if has_annexes %} are docxtpl control lines: they choose between
    two versions of a paragraph and disappear from the finished contract.

How it fits into the project:
    Run it once after every change to the skeleton wording:
        python templates/make_template.py
    It overwrites templates/contract_template.docx. The legal wording of the fixed
    parts lives here; the clause texts live in data/clauses.yaml.

Which documents it implements:
    - Contract skeleton (sections 1-13), copied word for word
    - Law packet Task 12 (title "Construction Contract"), Task 13 (prototype note
      below the title and in the footer), Task 14 (marker for sections 6-9),
      consistency c (comma in the section 7 heading)
    - Law colleague's answers 2 (texts of sections 2, 3 and 13 when there are no
      annexes), 3 (section 11 text while section 7 is empty), 4 (section 6: marker only)
    - Decisions D30 (one slot per clause group), D33 (all 13 sections stay)

Gaps that generator.py must fill:
    contract_date, client_name, client_address, contractor_name, contractor_address,
    site_address, work_description, exclusions_and_client_supplies, annex_register,
    has_annexes (true / false), price_clause, payment_clause,
    contractor_signatory_name, contractor_signatory_capacity
"""

from pathlib import Path  # builds the file path next to this script

from docx import Document  # python-docx: creates Word files


TEMPLATE_PATH = Path(__file__).parent / "contract_template.docx"  # where the template is saved

# marker for a section whose topic is not built yet (law packet Task 14, law colleague's answer 4)
NOT_GENERATED = "Not generated in this prototype. This section is incomplete; no waiver of statutory rights is intended."

# footer on every page (law packet Task 13)
FOOTER_TEXT = "Academic prototype - not for signature"


# ---------------------------------------------------------------------------
# Small helpers that add one piece to the document
# ---------------------------------------------------------------------------

def add_text(doc, text, italic=False):
    """Add one paragraph with the given text (optionally in italics)."""
    paragraph = doc.add_paragraph()  # a new, empty paragraph
    run = paragraph.add_run(text)  # the text as one piece ("run"), so a {{ gap }} is never split
    run.italic = italic  # italics on or off


def add_lines(doc, *lines):
    """Add one paragraph whose lines are separated by line breaks (like "Client:" / "Address:")."""
    paragraph = doc.add_paragraph()  # a new, empty paragraph
    for number, line in enumerate(lines):  # each line, counted from 0
        if number > 0:  # every line except the first ...
            paragraph.add_run().add_break()  # ... starts on a new line
        paragraph.add_run(line)  # the text of the line


def add_control(doc, tag):
    """Add a docxtpl control line such as {%p if has_annexes %}; it is removed when the contract is generated."""
    doc.add_paragraph(tag)  # must stand alone in its own paragraph (docxtpl rule for "{%p ... %}")


# ---------------------------------------------------------------------------
# The template, section by section
# ---------------------------------------------------------------------------

def build_template():
    """Build the template document and save it to templates/contract_template.docx."""
    doc = Document()  # a new, empty Word document

    # --- title block ---------------------------------------------------------
    doc.add_heading("Construction Contract", level=0)  # law packet Task 12 (skeleton said "iConstruction")
    add_text(doc, "Contract for residential building works in Switzerland", italic=True)  # skeleton subtitle
    add_text(doc, "Academic prototype for seminar demonstration. This document may contain incomplete sections "
                  "and requires completion and legal review before use.", italic=True)  # law packet Task 13
    add_text(doc, "Contract date: {{ contract_date }}")  # skeleton; formatted like "1 November 2026"

    # --- 1 Parties -------------------------------------------------------------
    doc.add_heading("1 Parties", level=2)  # skeleton section 1
    add_lines(doc, "Client: {{ client_name }}", "Address: {{ client_address }}")  # the Client
    add_lines(doc, "Contractor: {{ contractor_name }}", "Address: {{ contractor_address }}")  # the Contractor
    add_text(doc, "The Client and the Contractor are together referred to as the Parties.")  # skeleton

    # --- 2 Works and site ------------------------------------------------------
    doc.add_heading("2 Works and site", level=2)  # skeleton section 2
    add_control(doc, "{%p if has_annexes %}")  # version with technical annexes
    add_text(doc, "The Contractor shall produce and deliver the result described below and in the listed technical "
                  "annexes (the Works). The Client shall pay the remuneration determined under section 5.")  # skeleton
    add_control(doc, "{%p else %}")  # version without annexes
    add_text(doc, "The Contractor shall produce and deliver the result described below (the Works). The Client shall "
                  "pay the remuneration determined under section 5.")  # law colleague's answer 2
    add_control(doc, "{%p endif %}")  # end of the two versions
    add_text(doc, "Site address: {{ site_address }}")  # skeleton
    add_lines(doc, "Description of the Works and agreed result:", "{{ work_description }}")  # skeleton
    add_lines(doc, "Expressly excluded work and items supplied by the Client:",  # skeleton
              "{{ exclusions_and_client_supplies }}")  # empty answer -> text from parameters.yaml "if_empty_print"

    # --- 3 Contract documents --------------------------------------------------
    doc.add_heading("3 Contract documents", level=2)  # skeleton section 3
    add_control(doc, "{%p if has_annexes %}")  # version with technical annexes (skeleton)
    add_text(doc, "This Contract includes the technical annexes identified below by title, date and revision. "
                  "The Parties confirm that the listed annexes are available to both of them.")  # skeleton
    add_text(doc, "{{ annex_register }}")  # the list of annexes, one per line
    add_text(doc, "SIA 118 and either Party's standard terms are not incorporated into this Contract. Technical "
                  "annexes describe the agreed Works; they are not used to introduce additional general legal "
                  "conditions. A discrepancy affecting the Works shall be raised promptly and resolved before the "
                  "affected work proceeds where practicable. Statutory rights remain reserved.")  # skeleton
    add_control(doc, "{%p else %}")  # version without annexes: the whole body is replaced (answer 2)
    add_text(doc, "No technical annexes are listed or incorporated into this Contract.")  # law colleague's answer 2
    add_text(doc, "SIA 118 and either Party’s standard terms are not incorporated into this Contract. Any "
                  "discrepancy affecting the description of the Works shall be raised promptly and resolved before "
                  "the affected work proceeds where practicable. Statutory rights remain reserved.")  # answer 2
    add_control(doc, "{%p endif %}")  # end of the two versions

    # --- 4 Definitions ---------------------------------------------------------
    doc.add_heading("4 Definitions", level=2)  # skeleton section 4
    add_text(doc, "Delivery means making the completed Works available to the Client as required by the Contract. "
                  "Inspection means examining the Works for conformity and defects. Acceptance and approval have "
                  "their respective effects under applicable Swiss law. A delivery record or inspection protocol "
                  "shall identify defects and reservations; these concepts are not interchangeable.")  # skeleton

    # --- Contract terms (sections 5-10) ----------------------------------------
    doc.add_heading("Contract terms", level=1)  # skeleton part heading
    doc.add_heading("5 Price and payment", level=2)  # skeleton section 5 (topic 1)
    add_text(doc, "{{ price_clause }}")  # slot: PRICE_FIXED or PRICE_ESTIMATE (clauses.yaml "slot")
    add_text(doc, "{{ payment_clause }}")  # slot: PAYMENT_ON_COMPLETION or PAYMENT_MILESTONE

    # sections 6-9 belong to topics 2-5, which are not built yet: heading + marker only
    # (D33, law packet Task 14; for section 6 also the law colleague's answer 4)
    doc.add_heading("6 Completion date and delay penalty", level=2)  # topic 2 (later: deadline/penalty slots)
    add_text(doc, NOT_GENERATED)  # marker
    doc.add_heading("7 Inspection, acceptance and defects", level=2)  # topic 3; comma added (consistency c)
    add_text(doc, NOT_GENERATED)  # marker
    doc.add_heading("8 Liability", level=2)  # topic 4
    add_text(doc, NOT_GENERATED)  # marker
    doc.add_heading("9 Termination", level=2)  # topic 5
    add_text(doc, NOT_GENERATED)  # marker

    doc.add_heading("10 Changes to the Works", level=2)  # skeleton section 10 (PRICE_FIXED refers to it)
    add_text(doc, "The Parties shall record agreed changes to the Works and their effects on price and time in "
                  "writing before the changed work begins where practicable. A request or proposal does not itself "
                  "record agreement to a change. This procedure does not remove statutory duties to warn, statutory "
                  "relief or mandatory rights.")  # skeleton

    # --- General provisions and signatures (sections 11-13) --------------------
    doc.add_heading("General provisions and signatures", level=1)  # skeleton part heading
    doc.add_heading("11 Communications", level=2)  # skeleton section 11
    # whole body replaced while section 7 has no defect notice contact (law colleague's answer 3)
    add_text(doc, "Each Party shall promptly inform the other of changes to its contact details. Contract "
                  "communications may be sent to the relevant Party’s address stated in section 1 or to updated "
                  "contact details notified by that Party. Defect notices may be sent to the Contractor’s address "
                  "stated in section 1 or to updated defect-notice contact details notified by the Contractor. "
                  "Evidence of important communications should be kept. This section does not impose an additional "
                  "validity formality on a statutory notice; any applicable statutory form requirement remains "
                  "reserved.")  # answer 3

    doc.add_heading("12 Applicable law and disputes", level=2)  # skeleton section 12
    add_text(doc, "This Contract is governed by Swiss substantive law, including the applicable provisions on "
                  "contracts for work in the Swiss Code of Obligations.")  # skeleton, paragraph 1
    add_text(doc, "The Parties may first seek to resolve a dispute by direct discussion. Discussion is not a "
                  "condition for commencing proceedings or taking a step needed to preserve a right or "
                  "deadline.")  # skeleton, paragraph 2
    add_text(doc, "The courts having jurisdiction under applicable law shall have jurisdiction. No exclusive place "
                  "of jurisdiction or arbitration is agreed by this clause. Any mandatory jurisdiction and consumer "
                  "protections remain unaffected.")  # skeleton, paragraph 3

    doc.add_heading("13 Signatures", level=2)  # skeleton section 13
    add_control(doc, "{%p if has_annexes %}")  # version with technical annexes (skeleton)
    add_text(doc, "The Parties confirm that they have reviewed the completed Contract and the technical annexes "
                  "identified in section 3.")  # skeleton
    add_control(doc, "{%p else %}")  # version without annexes
    add_text(doc, "The Parties confirm that they have reviewed this Contract.")  # law colleague's answer 2
    add_control(doc, "{%p endif %}")  # end of the two versions
    add_lines(doc, "Client", "Name: {{ client_name }}",  # skeleton: the Client signs personally
              "Place and date: ______________________________", "Signature: __________________________________")  # lines
    add_lines(doc, "For the Contractor", "Name: {{ contractor_signatory_name }}",  # skeleton: someone signs for the company
              "Capacity: {{ contractor_signatory_capacity }}",  # e.g. "Managing director"
              "Place and date: ______________________________", "Signature: __________________________________")  # lines

    # --- footer on every page (law packet Task 13) -----------------------------
    doc.sections[0].footer.paragraphs[0].text = FOOTER_TEXT  # the document has one section, so this covers every page

    doc.save(TEMPLATE_PATH)  # write the Word file


if __name__ == "__main__":  # only runs when started directly: python templates/make_template.py
    build_template()  # build and save the template
    print(f"Template created: {TEMPLATE_PATH}")  # tell the user where it is
